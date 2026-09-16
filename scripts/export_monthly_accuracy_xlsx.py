#!/usr/bin/env python3
"""Export monthly SWIFT-TEC accuracy to Excel, including grid-cell metrics.

Two independent scores are produced:

1) Forecast accuracy
   Observed TEC vs reconstructed operational forecast using archived NOAA Kp
   forecast slots. The final forecast includes the existing polynomial Kp term
   plus the new Kp-band residual AI.

2) Kp-formula accuracy
   Observed 24h TEC change vs the existing polynomial Kp response change:
       dTEC_obs = TEC(t) - TEC(t-24h)
       dTEC_Kp  = F(Kp_actual(t)) - F(Kp_actual(t-24h))
   This deliberately excludes band residual AI so the polynomial itself can be
   judged separately.

Hit thresholds are exactly +/-5, +/-10, +/-15 TECU.

Sheets:
  Summary             overall + Kp-band metrics
  By_Lead             forecast metrics by +1/+2/+3/+4 day lead
  Forecast_Grid       per-grid, per-Kp-band forecast metrics
  KpFormula_Grid      per-grid, per-Kp-band formula metrics
  BandAI              learned band residual per grid cell
  Meta                coverage and method notes

The workbook is intended as a report artifact. It does not need to be served by
GitHub Pages and should normally NOT be committed daily, avoiding Git bloat.
"""
from __future__ import annotations

import argparse
import gzip
import json
import math
from bisect import bisect_left
from collections import defaultdict, OrderedDict
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

UTC = timezone.utc
THRESHOLDS = (5.0, 10.0, 15.0)
BAND_ORDER = ("ALL", "0-4", "5-6", "7-8", "9")
LEADS = (1, 2, 3, 4)


@dataclass(frozen=True)
class SourceCfg:
    name: str
    tec_root: Path
    ai_root: Path
    kp_archive: Path
    model: str


SOURCES = {
    "noaa": SourceCfg(
        name="noaa",
        tec_root=Path("docs/data/tec"),
        ai_root=Path("docs/data/ai"),
        kp_archive=Path("docs/data/ai/kp_actual_archive_recent.json"),
        model="prev_base",
    ),
    "isee": SourceCfg(
        name="isee",
        tec_root=Path("docs/data/isee_tec"),
        ai_root=Path("docs/data/ai/isee_japan"),
        kp_archive=Path("docs/data/ai/isee_japan/kp_actual_archive.json"),
        model="same_utc_mean",
    ),
}

FORECAST_ARCHIVE = Path("docs/data/ai/kp_forecast_archive.json")
ISEE_MEAN_DAYS = 10
STEP_MIN = 30
FORECAST_KP_TOL_MIN = 120


def iso(t):
    return t.astimezone(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def parse_time(s):
    try:
        return datetime.fromisoformat(str(s).replace("Z", "+00:00")).astimezone(UTC)
    except Exception:
        return None


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def load_json_maybe_gz(path: Path):
    if path.suffix == ".gz":
        with gzip.open(path, "rt", encoding="utf-8") as f:
            return json.load(f)
    return load_json(path)


def band_label(kp):
    x = float(kp)
    if x < 5.0:
        return "0-4"
    if x < 7.0:
        return "5-6"
    if x < 9.0:
        return "7-8"
    return "9"


def load_kp_rows(path: Path):
    if not path.exists():
        return []
    doc = load_json(path)
    rows = []
    for r in doc.get("rows", []):
        t = parse_time(r.get("time_utc") or r.get("time"))
        try:
            kp = float(r.get("kp"))
        except Exception:
            continue
        if t and math.isfinite(kp):
            rows.append((t, kp))
    rows.sort(key=lambda x: x[0])
    return rows


def nearest_kp(t, rows, max_hours=2.0):
    if not rows:
        return None
    times = [x[0] for x in rows]
    p = bisect_left(times, t)
    cand = []
    if p < len(rows):
        cand.append(rows[p])
    if p > 0:
        cand.append(rows[p - 1])
    if not cand:
        return None
    rt, kp = min(cand, key=lambda x: abs((x[0] - t).total_seconds()))
    if abs((rt - t).total_seconds()) > max_hours * 3600:
        return None
    return float(kp)


def load_metas(cfg: SourceCfg):
    p = cfg.tec_root / "index.json"
    if not p.exists():
        return []
    doc = load_json(p)
    out = []
    for r in doc.get("frames", []):
        t = parse_time(r.get("time_utc") or r.get("time"))
        rel = r.get("file") or r.get("path")
        if not t or not rel or t.minute % STEP_MIN != 0:
            continue
        path = cfg.tec_root / rel
        if path.exists():
            out.append({"time": t, "path": path, "file": rel})
    out.sort(key=lambda x: x["time"])
    return out


def nearest_meta(metas, target, tol_min=45):
    if not metas:
        return None
    times = [x["time"] for x in metas]
    p = bisect_left(times, target)
    cand = []
    if p < len(metas):
        cand.append(metas[p])
    if p > 0:
        cand.append(metas[p - 1])
    if not cand:
        return None
    best = min(cand, key=lambda x: abs((x["time"] - target).total_seconds()))
    if abs((best["time"] - target).total_seconds()) > tol_min * 60:
        return None
    return best


class FrameCache:
    def __init__(self, max_items=300):
        self.max_items = max_items
        self.d = OrderedDict()

    def get(self, meta):
        k = str(meta["path"])
        if k in self.d:
            v = self.d.pop(k)
            self.d[k] = v
            return v
        doc = load_json_maybe_gz(meta["path"])
        g = np.asarray(doc.get("grid") or [], dtype=float)
        if g.ndim != 2:
            raise ValueError(f"invalid grid: {meta['path']}")
        self.d[k] = (doc, g)
        while len(self.d) > self.max_items:
            self.d.popitem(last=False)
        return doc, g


class GridAgg:
    def __init__(self, shape):
        self.n = np.zeros(shape, dtype=np.int64)
        self.sum = np.zeros(shape, dtype=np.float64)
        self.abs = np.zeros(shape, dtype=np.float64)
        self.sq = np.zeros(shape, dtype=np.float64)
        self.hits = {th: np.zeros(shape, dtype=np.int64) for th in THRESHOLDS}

    def add(self, err, valid=None):
        e = np.asarray(err, dtype=float)
        mask = np.isfinite(e)
        if valid is not None:
            mask &= np.asarray(valid, dtype=bool)
        if not np.any(mask):
            return
        self.n[mask] += 1
        self.sum[mask] += e[mask]
        self.abs[mask] += np.abs(e[mask])
        self.sq[mask] += e[mask] * e[mask]
        for th in THRESHOLDS:
            self.hits[th][mask] += (np.abs(e[mask]) <= th).astype(np.int64)

    def scalar(self):
        n = int(self.n.sum())
        if n <= 0:
            return {"N": 0, "Hit5": None, "Hit10": None, "Hit15": None, "MAE": None, "RMSE": None, "Bias": None}
        s = float(self.sum.sum())
        ab = float(self.abs.sum())
        sq = float(self.sq.sum())
        return {
            "N": n,
            "Hit5": float(self.hits[5.0].sum()) / n,
            "Hit10": float(self.hits[10.0].sum()) / n,
            "Hit15": float(self.hits[15.0].sum()) / n,
            "MAE": ab / n,
            "RMSE": math.sqrt(sq / n),
            "Bias": s / n,
        }

    def cell_metrics(self, i, j):
        n = int(self.n[i, j])
        if n <= 0:
            return (0, None, None, None, None, None, None)
        return (
            n,
            self.hits[5.0][i, j] / n,
            self.hits[10.0][i, j] / n,
            self.hits[15.0][i, j] / n,
            self.abs[i, j] / n,
            math.sqrt(self.sq[i, j] / n),
            self.sum[i, j] / n,
        )


class ScalarAgg:
    def __init__(self):
        self.n = 0
        self.sum = 0.0
        self.abs = 0.0
        self.sq = 0.0
        self.hits = {th: 0 for th in THRESHOLDS}

    def add_array(self, err, valid=None):
        e = np.asarray(err, dtype=float)
        mask = np.isfinite(e)
        if valid is not None:
            mask &= np.asarray(valid, dtype=bool)
        if not np.any(mask):
            return
        x = e[mask]
        self.n += int(x.size)
        self.sum += float(x.sum())
        self.abs += float(np.abs(x).sum())
        self.sq += float(np.square(x).sum())
        for th in THRESHOLDS:
            self.hits[th] += int(np.count_nonzero(np.abs(x) <= th))

    def summary(self):
        if self.n <= 0:
            return {"N": 0, "Hit5": None, "Hit10": None, "Hit15": None, "MAE": None, "RMSE": None, "Bias": None}
        return {
            "N": self.n,
            "Hit5": self.hits[5.0] / self.n,
            "Hit10": self.hits[10.0] / self.n,
            "Hit15": self.hits[15.0] / self.n,
            "MAE": self.abs / self.n,
            "RMSE": math.sqrt(self.sq / self.n),
            "Bias": self.sum / self.n,
        }


def coeff_arrays(doc, month, shape):
    root = doc.get("coefficients_grid") or doc.get("grid_coefficients") or {}
    mg = root.get(str(month)) or root.get(month) or {}
    out = []
    for k in ("k0", "k1", "k2", "k3"):
        try:
            a = np.asarray(mg.get(k), dtype=float)
            if a.shape != shape:
                a = np.zeros(shape)
        except Exception:
            a = np.zeros(shape)
        out.append(np.nan_to_num(a, nan=0.0, posinf=0.0, neginf=0.0))
    return out


def F_grid(coeff_doc, month, kp, shape):
    k0, k1, k2, k3 = coeff_arrays(coeff_doc, month, shape)
    x = float(kp) - 3.0
    v = k0 + k1 * x + k2 * x * x + k3 * x * x * x
    clip = float(coeff_doc.get("correction_clip_tecu") or 20.0)
    return np.clip(np.nan_to_num(v, nan=0.0), -clip, clip)


def load_band_doc(cfg, month):
    p = cfg.ai_root / "kp_band_residual" / f"{month:02d}.json"
    if not p.exists():
        return {}
    return load_json(p)


def band_fallback_order(label):
    if label == "9":
        return ("9", "7-8", "5-6", "0-4")
    if label == "7-8":
        return ("7-8", "5-6", "0-4")
    if label == "5-6":
        return ("5-6", "0-4")
    return ("0-4",)


def band_residual_grid(doc, kp, shape):
    if not doc:
        return np.zeros(shape, dtype=float)
    label = band_label(kp)
    bands = doc.get("bands") or {}
    for bname in band_fallback_order(label):
        b = bands.get(bname) or {}
        try:
            r = np.asarray(b.get("residual"), dtype=float)
            n = np.asarray(b.get("sample_count"), dtype=float)
            if r.shape != shape or n.shape != shape:
                continue
            min_apply = int(b.get("min_apply_samples") or 1)
            mask = np.isfinite(r) & (n >= min_apply)
            if np.any(mask):
                # Fallback is cell-wise: cells without enough data fall through to
                # the next band. Build a result recursively below.
                result = np.full(shape, np.nan, dtype=float)
                result[mask] = r[mask]
                # fill missing cells from lower bands
                for lower in band_fallback_order(label)[band_fallback_order(label).index(bname)+1:]:
                    lb = bands.get(lower) or {}
                    try:
                        lr = np.asarray(lb.get("residual"), dtype=float)
                        ln = np.asarray(lb.get("sample_count"), dtype=float)
                        if lr.shape != shape or ln.shape != shape:
                            continue
                        lm = np.isfinite(lr) & (ln >= int(lb.get("min_apply_samples") or 1)) & ~np.isfinite(result)
                        result[lm] = lr[lm]
                    except Exception:
                        pass
                return np.nan_to_num(result, nan=0.0)
        except Exception:
            continue
    return np.zeros(shape, dtype=float)


def load_real_forecast_slots():
    if not FORECAST_ARCHIVE.exists():
        return {lead: [] for lead in LEADS}
    doc = load_json(FORECAST_ARCHIVE)
    by_lead = {lead: [] for lead in LEADS}
    for issue_doc in doc.get("forecasts", []):
        issue = parse_time(issue_doc.get("issue_utc"))
        src_type = str(issue_doc.get("source_type") or issue_doc.get("source_url") or "")
        if "cold_start" in src_type:
            continue
        for s in issue_doc.get("slots", []):
            t = parse_time(s.get("time_utc"))
            try:
                kp = float(s.get("kp_forecast"))
            except Exception:
                continue
            if not t or not math.isfinite(kp):
                continue
            lead = int(s.get("lead_day") or (math.ceil(max(0, (t - issue).total_seconds()) / 86400.0) if issue else 0))
            if lead not in LEADS:
                continue
            src = str(s.get("forecast_source") or src_type)
            if "cold_start" in src:
                continue
            by_lead[lead].append({"time": t, "kp": kp, "issue": issue, "source": src})
    for lead in LEADS:
        by_lead[lead].sort(key=lambda x: (x["time"], x["issue"] or datetime.min.replace(tzinfo=UTC)))
    return by_lead


def forecast_slot_for_target(target, lead, by_lead):
    rows = by_lead.get(lead) or []
    if not rows:
        return None
    times = [r["time"] for r in rows]
    p = bisect_left(times, target)
    cand = []
    # Multiple issues can map to the same target slot; inspect a local window.
    for k in range(max(0, p - 8), min(len(rows), p + 9)):
        r = rows[k]
        if abs((r["time"] - target).total_seconds()) <= FORECAST_KP_TOL_MIN * 60:
            cand.append(r)
    if not cand:
        return None
    # Prefer closest target time, then the latest issue that still belongs to
    # this lead-day bucket.
    cand.sort(key=lambda r: (abs((r["time"] - target).total_seconds()), -(r["issue"].timestamp() if r["issue"] else 0)))
    return cand[0]


def weighted_isee_mean_before(seq, issue, cache, shape):
    prior = [m for m in seq if m["time"] < issue]
    if not prior:
        return None
    hist = prior[-ISEE_MEAN_DAYS:]
    num = np.zeros(shape, dtype=float)
    den = np.zeros(shape, dtype=float)
    for rank, meta in enumerate(reversed(hist)):
        try:
            _, g = cache.get(meta)
        except Exception:
            continue
        if g.shape != shape:
            continue
        w = max(0.1, 1.0 - 0.1 * rank)
        mask = np.isfinite(g)
        num[mask] += g[mask] * w
        den[mask] += w
    if not np.any(den > 0):
        return None
    with np.errstate(invalid="ignore", divide="ignore"):
        return np.where(den > 0, num / den, np.nan)


def month_bounds(month_str):
    y, m = map(int, month_str.split("-"))
    start = datetime(y, m, 1, tzinfo=UTC)
    if m == 12:
        end = datetime(y + 1, 1, 1, tzinfo=UTC)
    else:
        end = datetime(y, m + 1, 1, tzinfo=UTC)
    return start, end


def infer_month(metas):
    if not metas:
        return datetime.now(UTC).strftime("%Y-%m")
    return metas[-1]["time"].strftime("%Y-%m")


def score_source(cfg: SourceCfg, month_str: str):
    metas = load_metas(cfg)
    if not metas:
        raise RuntimeError(f"{cfg.name}: no TEC frames")
    kp_rows = load_kp_rows(cfg.kp_archive)
    if not kp_rows:
        raise RuntimeError(f"{cfg.name}: Kp actual archive empty: {cfg.kp_archive}")
    coeff_path = cfg.ai_root / "kp_grid_coefficients.json"
    if not coeff_path.exists():
        raise RuntimeError(f"{cfg.name}: missing {coeff_path}")
    coeff_doc = load_json(coeff_path)
    forecast_slots = load_real_forecast_slots()
    cache = FrameCache(max_items=360)

    first_doc, first_grid = cache.get(metas[0])
    shape = first_grid.shape
    lat_arr = list(first_doc.get("lat_arr") or first_doc.get("latArr") or range(shape[0]))
    lon_arr = list(first_doc.get("lon_arr") or first_doc.get("lonArr") or range(shape[1]))
    if len(lat_arr) != shape[0] or len(lon_arr) != shape[1]:
        raise RuntimeError(f"{cfg.name}: lat/lon size mismatch")

    formula_aggs = {b: GridAgg(shape) for b in BAND_ORDER}
    forecast_aggs = {b: GridAgg(shape) for b in BAND_ORDER}
    lead_aggs = {lead: ScalarAgg() for lead in LEADS}

    start, end = month_bounds(month_str)
    targets = [m for m in metas if start <= m["time"] < end]

    # ISEE same-UTC-time lookup for reconstructing forecasts at each archived issue.
    by_slot = defaultdict(list)
    if cfg.name == "isee":
        for m in metas:
            by_slot[m["time"].strftime("%H%M")].append(m)
        for k in by_slot:
            by_slot[k].sort(key=lambda x: x["time"])

    formula_frames = 0
    forecast_cases = 0
    forecast_slots_missing = 0

    for mf in targets:
        t = mf["time"]
        try:
            _, actual = cache.get(mf)
        except Exception:
            continue
        if actual.shape != shape:
            continue

        kp_actual = nearest_kp(t, kp_rows)
        if kp_actual is None:
            continue
        actual_band = band_label(kp_actual)

        # --- Kp formula accuracy: isolate 24h Kp-response equation ---
        mb24 = nearest_meta(metas, t - timedelta(hours=24), tol_min=45)
        if mb24:
            kp_b = nearest_kp(mb24["time"], kp_rows)
            if kp_b is not None:
                try:
                    _, base24 = cache.get(mb24)
                except Exception:
                    base24 = None
                if base24 is not None and base24.shape == shape:
                    obs_delta = actual - base24
                    pred_delta = F_grid(coeff_doc, t.month, kp_actual, shape) - F_grid(coeff_doc, mb24["time"].month, kp_b, shape)
                    ferr = obs_delta - pred_delta
                    valid = np.isfinite(actual) & np.isfinite(base24)
                    formula_aggs["ALL"].add(ferr, valid)
                    formula_aggs[actual_band].add(ferr, valid)
                    formula_frames += 1

        # --- Final forecast accuracy using real archived Kp forecasts ---
        for lead in LEADS:
            fs = forecast_slot_for_target(t, lead, forecast_slots)
            if not fs:
                forecast_slots_missing += 1
                continue
            kp_f = float(fs["kp"])
            issue = fs["issue"]
            if not issue:
                continue

            if cfg.name == "noaa":
                mb = nearest_meta(metas, t - timedelta(days=lead), tol_min=90)
                if not mb:
                    continue
                kp_b = nearest_kp(mb["time"], kp_rows)
                if kp_b is None:
                    continue
                try:
                    _, base = cache.get(mb)
                except Exception:
                    continue
                if base.shape != shape:
                    continue
                corr = F_grid(coeff_doc, t.month, kp_f, shape) - F_grid(coeff_doc, mb["time"].month, kp_b, shape)
                band_doc = load_band_doc(cfg, t.month)
                corr += band_residual_grid(band_doc, kp_f, shape)
                clip = float(coeff_doc.get("correction_clip_tecu") or 20.0)
                corr = np.clip(corr, -clip, clip)
                forecast = base + corr
                valid = np.isfinite(actual) & np.isfinite(base)
            else:
                seq = by_slot.get(t.strftime("%H%M")) or []
                base = weighted_isee_mean_before(seq, issue, cache, shape)
                if base is None:
                    continue
                corr = F_grid(coeff_doc, t.month, kp_f, shape)
                band_doc = load_band_doc(cfg, t.month)
                corr += band_residual_grid(band_doc, kp_f, shape)
                clip = float(coeff_doc.get("correction_clip_tecu") or 20.0)
                corr = np.clip(corr, -clip, clip)
                forecast = base + corr
                valid = np.isfinite(actual) & np.isfinite(base)

            forecast = np.clip(forecast, 0.0, 300.0)
            err = actual - forecast
            forecast_aggs["ALL"].add(err, valid)
            forecast_aggs[actual_band].add(err, valid)
            lead_aggs[lead].add_array(err, valid)
            forecast_cases += 1

    meta = {
        "source": cfg.name,
        "month": month_str,
        "coverage_first": iso(targets[0]["time"]) if targets else None,
        "coverage_last": iso(targets[-1]["time"]) if targets else None,
        "tec_frames_in_month": len(targets),
        "formula_frames_scored": formula_frames,
        "forecast_target_lead_cases": forecast_cases,
        "forecast_slot_misses": forecast_slots_missing,
        "grid_shape": shape,
        "forecast_note": (
            "Uses real archived NOAA Kp forecast slots only; cold-start pseudo forecasts are excluded. "
            "Historical forecast grids are reconstructed with the currently stored month coefficients/band AI."
        ),
        "formula_note": "Kp formula score compares observed 24h TEC change with F(Kp_actual_t)-F(Kp_actual_t-24h); band AI excluded.",
    }
    return {
        "meta": meta,
        "lat_arr": lat_arr,
        "lon_arr": lon_arr,
        "formula": formula_aggs,
        "forecast": forecast_aggs,
        "lead": lead_aggs,
        "band_doc": load_band_doc(cfg, int(month_str.split("-")[1])),
    }


def pct(v):
    return None if v is None else float(v)


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[row]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")


def autofit(ws, max_width=28):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 10
        for c in col[:300]:
            if c.value is not None:
                width = max(width, min(max_width, len(str(c.value)) + 2))
        ws.column_dimensions[letter].width = width


def write_summary(wb, result):
    ws = wb.create_sheet("Summary")
    ws.append(["Metric", "KpBand", "N", "Hit ±5", "Hit ±10", "Hit ±15", "MAE TECU", "RMSE TECU", "Bias TECU"])
    for metric_name, aggs in (("Forecast", result["forecast"]), ("KpFormula", result["formula"])):
        for band in BAND_ORDER:
            s = aggs[band].scalar()
            ws.append([metric_name, band, s["N"], pct(s["Hit5"]), pct(s["Hit10"]), pct(s["Hit15"]), s["MAE"], s["RMSE"], s["Bias"]])
    style_header(ws)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for c in row:
            c.number_format = "0.0%"
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=9):
        for c in row:
            c.number_format = "0.000"
    ws.freeze_panes = "A2"
    ws.conditional_formatting.add(f"D2:F{ws.max_row}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    autofit(ws)


def write_lead(wb, result):
    ws = wb.create_sheet("By_Lead")
    ws.append(["LeadDay", "N", "Hit ±5", "Hit ±10", "Hit ±15", "MAE TECU", "RMSE TECU", "Bias TECU"])
    for lead in LEADS:
        s = result["lead"][lead].summary()
        ws.append([lead, s["N"], s["Hit5"], s["Hit10"], s["Hit15"], s["MAE"], s["RMSE"], s["Bias"]])
    style_header(ws)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for c in row:
            c.number_format = "0.0%"
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=8):
        for c in row:
            c.number_format = "0.000"
    autofit(ws)


def write_grid_sheet(wb, name, aggs, lat_arr, lon_arr):
    ws = wb.create_sheet(name)
    ws.append(["KpBand", "Lat", "Lon", "N", "Hit ±5", "Hit ±10", "Hit ±15", "MAE TECU", "RMSE TECU", "Bias TECU"])
    for band in BAND_ORDER:
        a = aggs[band]
        for i, lat in enumerate(lat_arr):
            for j, lon in enumerate(lon_arr):
                n, h5, h10, h15, mae, rmse, bias = a.cell_metrics(i, j)
                if n <= 0:
                    continue
                ws.append([band, float(lat), float(lon), n, h5, h10, h15, mae, rmse, bias])
    style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=7):
        for c in row:
            c.number_format = "0.0%"
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=10):
        for c in row:
            c.number_format = "0.000"
    if ws.max_row > 1:
        ws.conditional_formatting.add(f"E2:G{ws.max_row}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    autofit(ws)


def write_band_ai(wb, result):
    ws = wb.create_sheet("BandAI")
    ws.append(["KpBand", "Lat", "Lon", "Residual TECU", "SampleCount", "MinApply"])
    doc = result.get("band_doc") or {}
    bands = doc.get("bands") or {}
    lat_arr = result["lat_arr"]
    lon_arr = result["lon_arr"]
    for band in ("0-4", "5-6", "7-8", "9"):
        b = bands.get(band) or {}
        try:
            r = np.asarray(b.get("residual"), dtype=float)
            n = np.asarray(b.get("sample_count"), dtype=int)
        except Exception:
            continue
        if r.shape != (len(lat_arr), len(lon_arr)) or n.shape != r.shape:
            continue
        min_apply = int(b.get("min_apply_samples") or 0)
        for i, lat in enumerate(lat_arr):
            for j, lon in enumerate(lon_arr):
                if n[i, j] <= 0 and abs(r[i, j]) < 1e-12:
                    continue
                ws.append([band, float(lat), float(lon), float(r[i, j]), int(n[i, j]), min_apply])
    style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws["D"][1:]:
        c.number_format = "0.000"
    autofit(ws)


def write_meta(wb, result):
    ws = wb.create_sheet("Meta")
    ws.append(["Key", "Value"])
    for k, v in result["meta"].items():
        ws.append([k, str(v)])
    ws.append(["hit_thresholds_tecu", "5,10,15"])
    ws.append(["kp_bands", "0-4 = Kp<5; 5-6 = 5<=Kp<7; 7-8 = 7<=Kp<9; 9 = Kp>=9"])
    style_header(ws)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(result, output: Path):
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    write_summary(wb, result)
    write_lead(wb, result)
    write_grid_sheet(wb, "Forecast_Grid", result["forecast"], result["lat_arr"], result["lon_arr"])
    write_grid_sheet(wb, "KpFormula_Grid", result["formula"], result["lat_arr"], result["lon_arr"])
    write_band_ai(wb, result)
    write_meta(wb, result)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", choices=["noaa", "isee", "both"], default="both")
    ap.add_argument("--month", default="auto", help="YYYY-MM or auto (latest observation month per source)")
    ap.add_argument("--output-root", default="reports/monthly_accuracy")
    args = ap.parse_args()

    selected = ["noaa", "isee"] if args.source == "both" else [args.source]
    for name in selected:
        cfg = SOURCES[name]
        metas = load_metas(cfg)
        month = infer_month(metas) if args.month == "auto" else args.month
        print(f"=== {name}: scoring month {month} ===")
        result = score_source(cfg, month)
        out = Path(args.output_root) / name / f"{name}_{month}.xlsx"
        build_workbook(result, out)
        print(f"WROTE {out} / {result['meta']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
